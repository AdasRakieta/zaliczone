"""Build the illustrated report from the retained UKW DOCX template."""
import argparse
import copy
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image
from report_figures import generate

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/report/template.docx'
ASSETS = ROOT / 'docs/report/assets'
FINAL = ROOT / 'output/docx/Sprawozdanie-Zaliczone.docx'
PDF_DIR = ROOT / 'output/pdf'
WORK = ROOT / 'tmp/report-work'
spec = importlib.util.spec_from_file_location('report_content', ROOT/'docs/report/content.py')
content = importlib.util.module_from_spec(spec)
spec.loader.exec_module(content)
PAGES = content.PAGES
CROPS = json.loads((ASSETS/'crops.json').read_text(encoding='utf-8-sig'))
doc = None
figure_number = 0

def font(run, size=12, bold=None, italic=None):
    run.font.name = 'Times New Roman'; run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    run.font.color.rgb = RGBColor(0,0,0)
    return run

def paragraph(text, compact=False):
    item = doc.add_paragraph(style='Normal')
    item.paragraph_format.widow_control = True
    if compact:
        item.alignment = WD_ALIGN_PARAGRAPH.LEFT
        item.paragraph_format.first_line_indent = Cm(0)
        item.paragraph_format.line_spacing = 1.12
        item.paragraph_format.space_after = Pt(4)
    for i, line in enumerate(text.split('\n')):
        if i: item.add_run().add_break()
        font(item.add_run(line), 11.5 if compact else 12)
    return item

def heading(text, level=2, new_page=False):
    item = doc.add_paragraph(text, style='Heading 1' if level == 1 else 'Heading 2')
    item.paragraph_format.page_break_before = new_page
    item.paragraph_format.keep_with_next = True
    item.paragraph_format.space_before = Pt(0 if new_page else 8)
    item.paragraph_format.space_after = Pt(6)
    if level == 2:
        item.paragraph_format.left_indent = Cm(1.8)
        item.paragraph_format.first_line_indent = Cm(-1.15)
        item.paragraph_format.tab_stops.clear_all()
        item.paragraph_format.tab_stops.add_tab_stop(Cm(1.8))
    for run in item.runs: run.font.color.rgb = RGBColor(0,0,0)
    return item

def add_table(block):
    tab = doc.add_table(rows=1, cols=len(block['headers']))
    tab.autofit = False
    widths = block.get('widths') or [15.8/len(block['headers'])]*len(block['headers'])
    for column, width in zip(tab.columns, widths): column.width = Cm(width)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        line=OxmlElement('w:'+edge)
        for key,value in {'val':'single','sz':'4','color':'D9D9D9'}.items():line.set(qn('w:'+key),value)
        borders.append(line)
    tab._tbl.tblPr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','bottom','left','right']:
        part=OxmlElement('w:'+edge);part.set(qn('w:w'),'65' if edge in ['top','bottom'] else '85');part.set(qn('w:type'),'dxa');margins.append(part)
    tab._tbl.tblPr.append(margins)
    rows=[block['headers']]+block['rows']
    for row_index, values in enumerate(rows):
        row=tab.rows[0] if row_index==0 else tab.add_row()
        props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
        if row_index==0:props.append(OxmlElement('w:tblHeader'))
        for column_index,(cell,value,width) in enumerate(zip(row.cells,values,widths)):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para=cell.paragraphs[0];para.paragraph_format.first_line_indent=Cm(0)
            para.paragraph_format.line_spacing=1.06;para.paragraph_format.space_after=Pt(2)
            para.paragraph_format.space_before=Pt(2)
            para.alignment=WD_ALIGN_PARAGRAPH.CENTER if width<2.7 else WD_ALIGN_PARAGRAPH.LEFT
            emphasized = row_index == 0 or row_index in block.get('emphasis_rows',[])
            font(para.add_run(str(value)),10.5,bold=emphasized)
            if emphasized:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');cell._tc.get_or_add_tcPr().append(shade)
    gap=doc.add_paragraph();gap.paragraph_format.first_line_indent=Cm(0)
    gap.paragraph_format.line_spacing=1;gap.paragraph_format.space_after=Pt(0)
    font(gap.add_run(''),3)
    gap.paragraph_format.line_spacing=Pt(3)

def picture(run, filename, width, max_height):
    path=ASSETS/filename
    with Image.open(path) as image: original_width,original_height=image.size
    crop=CROPS.get(filename,{'x':0,'y':0,'width':original_width,'height':original_height})
    x,y=max(0,crop['x']),max(0,crop['y'])
    crop_width=min(crop['width'],original_width-x);crop_height=min(crop['height'],original_height-y)
    if crop_width <= 0 or crop_height <= 0:raise ValueError('Invalid image crop: '+filename)
    height=width*crop_height/crop_width
    if height>max_height:width*=max_height/height;height=max_height
    shape=run.add_picture(str(path),width=Cm(width),height=Cm(height))
    shape._inline.docPr.set('descr',filename.replace('.png','').replace('-',' '))
    rect=OxmlElement('a:srcRect')
    for key,value in {'l':x/original_width,'t':y/original_height,'r':(original_width-x-crop_width)/original_width,'b':(original_height-y-crop_height)/original_height}.items():rect.set(key,str(round(value*100000)))
    shape._inline.graphic.graphicData.pic.blipFill.insert(1,rect)

def add_figure(block):
    global figure_number
    figure_number+=1
    para=doc.add_paragraph();para.paragraph_format.first_line_indent=Cm(0)
    para.paragraph_format.line_spacing=1;para.paragraph_format.space_before=Pt(5)
    para.paragraph_format.space_after=Pt(0);para.paragraph_format.keep_with_next=True
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if block['type']=='pair':
        for i,(filename,width) in enumerate(zip(block['files'],block['widths'])):
            if i:font(para.add_run('  '),12)
            picture(para.add_run(),filename,width,block['max_height'])
    else:picture(para.add_run(),block['file'],block['width'],block['max_height'])
    caption=doc.add_paragraph();caption.paragraph_format.first_line_indent=Cm(0)
    caption.alignment=WD_ALIGN_PARAGRAPH.CENTER;caption.paragraph_format.line_spacing=1.1
    caption.paragraph_format.space_before=Pt(3);caption.paragraph_format.space_after=Pt(6)
    font(caption.add_run(f"Rys. {figure_number}. {block['caption']}"),10,italic=True)

def add_notes(block):
    title=paragraph(block['title'],True);title.paragraph_format.space_before=Pt(5)
    title.paragraph_format.keep_with_next=True
    for run in title.runs:run.bold=True
    for person,speech in block['entries']:
        item=doc.add_paragraph();item.paragraph_format.first_line_indent=Cm(0)
        item.alignment=WD_ALIGN_PARAGRAPH.LEFT
        item.paragraph_format.line_spacing=1.08;item.paragraph_format.space_after=Pt(2)
        font(item.add_run(person+': '),11.5,bold=True);font(item.add_run(speech),11.5)
    decision=doc.add_paragraph();decision.paragraph_format.first_line_indent=Cm(0)
    decision.alignment=WD_ALIGN_PARAGRAPH.LEFT
    decision.paragraph_format.line_spacing=1.08;decision.paragraph_format.space_after=Pt(4)
    font(decision.add_run('Ustalenie: '),11,bold=True);font(decision.add_run(block['decision']),11)

def build_toc():
    title=doc.add_paragraph('Spis treści',style='Normal')
    title.paragraph_format.first_line_indent=Cm(0);font(title.runs[0],18,bold=True)
    title.paragraph_format.space_after=Pt(12)
    chapter=0
    for page_number,page in enumerate(PAGES,3):
        if 'chapter' not in page:continue
        chapter+=1
        line=paragraph(f"{chapter}. {page['chapter']}\t{page_number}",True)
        line.paragraph_format.tab_stops.add_tab_stop(Cm(15.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        line.paragraph_format.space_after=Pt(5)
    sub=paragraph('Spis ilustracji',True)
    font(sub.runs[0],14,bold=True)
    number=0
    for page_number,page in enumerate(PAGES,3):
        for block in page['blocks']:
            if block['type'] not in ['image','pair']:continue
            number+=1
            line=paragraph(f"Rys. {number}. {block['caption']}\t{page_number}",True)
            for run in line.runs:run.font.size=Pt(10)
            line.paragraph_format.line_spacing=1.05
            line.paragraph_format.space_after=Pt(3)
            line.paragraph_format.tab_stops.add_tab_stop(Cm(15.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    paragraph('Temat opracowania: Zaliczone, planer nauki i zaliczeń dla studentów. Dokument zachowuje układ dziesięciu rozdziałów i logo z przekazanego wzorca UKW. Notatki spotkań, daty sprintów, estymaty i feedback 360° tworzą scenariusz dydaktyczny pracy czteroosobowego zespołu. Zrzuty pokazują aplikację i materiały projektu.',True)

def build_markdown():
    lines=['# Sprawozdanie projektu Zaliczone','', '**Zespół:** Szymon Przybysz, Iwona Figas, Jarek Antczak, Emil Kosicki.','',
           'Wersja składu zgodna z wzorcem UKW znajduje się w output/docx/Sprawozdanie-Zaliczone.docx i output/pdf/Sprawozdanie-Zaliczone.pdf. Notatki i feedback stanowią symulację dydaktyczną.','']
    chapter=0;image_index=0
    for page in PAGES:
        if 'chapter' in page:chapter+=1;lines += [f"## {chapter}. {page['chapter']}",'']
        else:lines += ['### '+page['title'],'']
        for block in page['blocks']:
            kind=block['type']
            if kind=='p':lines += [block['text'],'']
            elif kind=='h':lines += ['### '+block['text'],'']
            elif kind=='table':
                lines += ['| '+' | '.join(block['headers'])+' |','| '+' | '.join(['---']*len(block['headers']))+' |']
                lines += ['| '+' | '.join(str(value).replace('\n','<br>') for value in row)+' |' for row in block['rows']];lines+=['']
            elif kind in ['image','pair']:
                image_index+=1
                files=block['files'] if kind=='pair' else [block['file']]
                lines += [f"![{block['caption']}](report/assets/{filename})" for filename in files]
                lines += [f"Rys. {image_index}. {block['caption']}",'']
            elif kind=='notes':
                lines+=['#### '+block['title'],'']
                lines += [f"- **{person}:** {speech}" for person,speech in block['entries']]
                lines += ['', '**Ustalenie:** '+block['decision'],'']
            elif kind=='checks':lines += ['- [x] '+text for text in block['items']]+['']
    (ROOT/'docs/REPORT.md').write_text('\n'.join(lines),encoding='utf8')

def main():
    global doc
    parser=argparse.ArgumentParser();parser.add_argument('--pdf',action='store_true');args=parser.parse_args()
    WORK.mkdir(parents=True,exist_ok=True);FINAL.parent.mkdir(parents=True,exist_ok=True)
    source_hash=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    generate();doc=Document(SOURCE)
    body=doc._element.body;ending=copy.deepcopy(body.sectPr)
    for child in list(body)[14:]:body.remove(child)
    body.append(ending)
    cover=doc.paragraphs
    cover[7].clear();font(cover[7].add_run('Zaliczone'),16.5,bold=True)
    cover[11].clear();font(cover[11].add_run('Szymon Przybysz, Iwona Figas\nJarek Antczak, Emil Kosicki'),16.5,bold=True)
    cover[13].clear();font(cover[13].add_run('2026/2027'),16.5,bold=True)
    section=doc.sections[1];section.footer.is_linked_to_previous=False
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.first_line_indent=Cm(0)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    number=OxmlElement('w:pgNumType');number.set(qn('w:start'),'2')
    old_number=section._sectPr.find(qn('w:pgNumType'))
    if old_number is not None:section._sectPr.remove(old_number)
    section._sectPr.append(number)
    build_toc()
    for page in PAGES:
        heading(page.get('chapter',page.get('title')),1 if 'chapter' in page else 2,True)
        for block in page['blocks']:
            kind=block['type']
            if kind=='p':paragraph(block['text'],block.get('compact',False))
            elif kind=='h':heading(block['text'])
            elif kind=='table':add_table(block)
            elif kind in ['image','pair']:add_figure(block)
            elif kind=='notes':add_notes(block)
            elif kind=='checks':
                for text in block['items']:paragraph('[x] '+text,True)
    doc.core_properties.title='Sprawozdanie projektu Zaliczone'
    doc.core_properties.author='Szymon Przybysz; Iwona Figas; Jarek Antczak; Emil Kosicki'
    doc.core_properties.last_modified_by='Zespół Zaliczone'
    temporary=WORK/'report-working.docx';doc.save(temporary)
    mutable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml','docProps/app.xml'}
    with zipfile.ZipFile(SOURCE) as source,zipfile.ZipFile(temporary) as generated,zipfile.ZipFile(FINAL,'w',zipfile.ZIP_DEFLATED) as out:
        original=set(source.namelist());created=set(generated.namelist())
        for name in sorted(original|created):
            data=generated.read(name) if name in mutable or name not in original else source.read(name)
            out.writestr(name,data)
    with zipfile.ZipFile(SOURCE) as source,zipfile.ZipFile(FINAL) as final:
        preserved=[]
        for name in source.namelist():
            if name not in mutable:
                assert source.read(name)==final.read(name),f'Template part changed: {name}'
                preserved.append(name)
        original_rels=Document(SOURCE).part.rels
        final_rels=Document(FINAL).part.rels
        for rid,rel in original_rels.items():
            assert rid in final_rels and rel.target_ref==final_rels[rid].target_ref
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
    proof={'template_sha256':source_hash,'content_pages':len(PAGES),'expected_pages':len(PAGES)+2,'chapters':[page['chapter'] for page in PAGES if 'chapter' in page],'figures':figure_number,'preserved_parts':preserved}
    (WORK/'fidelity.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2),encoding='utf8')
    build_markdown()
    if args.pdf:
        executable=os.environ.get('LIBREOFFICE_PATH') or shutil.which('soffice.exe') or shutil.which('soffice')
        local=ROOT/'tmp/document-runtime/libreoffice/program/soffice.exe'
        if not executable and local.exists():executable=str(local)
        if not executable:raise RuntimeError('Set LIBREOFFICE_PATH to LibreOffice soffice executable.')
        PDF_DIR.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='zaliczone-lo-') as profile, tempfile.TemporaryDirectory(prefix='zaliczone-pdf-') as export:
            subprocess.run([executable,'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--norestore','--convert-to','pdf','--outdir',export,str(FINAL)],check=True)
            exported=Path(export)/(FINAL.stem+'.pdf')
            if not exported.is_file():raise RuntimeError('PDF export failed')
            import pymupdf
            with pymupdf.open(exported) as pdf:
                if len(pdf)!=len(PAGES)+2:raise RuntimeError(f'Expected {len(PAGES)+2} pages, got {len(pdf)}; update layout before publishing.')
                for index,page in enumerate(PAGES,2):
                    text=' '.join(pdf[index].get_text().split())
                    expected=page.get('chapter',page.get('title'))
                    if expected not in text:raise RuntimeError(f'TOC mismatch on page {index+1}: {expected}')
            shutil.copyfile(exported,PDF_DIR/exported.name)
    print(json.dumps(proof,ensure_ascii=True))

if __name__=='__main__':main()
